import os
import json
import time
import gradio as gr

MAX_SLIDES = 25
SESSIONS_DIR = "presentation_sessions"
os.makedirs(SESSIONS_DIR, exist_ok=True)

def process_outline_and_suggest_layouts(slide_count, outline_text):
    """Processes outline text data and returns layout recommendations for exactly 25 tabs."""
    count = int(slide_count)
    updates = [gr.update(visible=True)]  # Show workspace parent container
    lines = [line.strip() for line in outline_text.split("\n") if line.strip()]
    
    for i in range(MAX_SLIDES):
        if i < count:
            slide_title = lines[i] if i < len(lines) else f"Slide {i+1} Title"
            suggested_layout = "Split Text & Image"
            llm_insights = "💡 Focus on clear typography. Keep content below 4 bullets."
            llm_questions = "❓ Should we include a timeline graphic here?"
            
            if "Architecture" in slide_title or "Technical" in slide_title:
                suggested_layout = "Technical Diagram Grid"
                llm_insights = "💡 Best suited for system topologies or structural layouts."
            elif "Introduction" in slide_title or "Conclusion" in slide_title:
                suggested_layout = "Minimalist Title Centered"
            
            updates.extend([
                gr.update(visible=True),  # 1. slide_tab visibility
                gr.update(value=slide_title),  # 2. slide_title Textbox
                gr.update(value=f"• Key point or bullet draft 1\n• Key point or bullet draft 2"), # 3. slide_content
                gr.update(value=suggested_layout),  # 4. slide_layout text tracker
                gr.update(value=llm_insights),  # 5. llm_insights_box
                gr.update(value=llm_questions),  # 6. llm_questions_box
                gr.update(value=""),  # 7. user_feedback_box reset
                gr.update(value=None),  # 8. slide_preview reset
                gr.update(value=False)  # 9. approve_chk reset
            ])
        else:
            # Hide unused slide rows completely up to 25 to avoid parameter length drift
            updates.extend([gr.update(visible=False)] + [gr.update() for _ in range(8)])
            
    return updates


def save_user_session(session_name, count, outline, *slide_states):
    """Serializes the complete runtime state matrix into a localized tracking JSON block."""
    if not session_name.strip():
        return "❌ Please enter a valid Session Name identifier first!"
        
    slides_data = []
    step = 8  # Number of saved element values mapped inside state_components_to_persist
    
    for i in range(MAX_SLIDES):
        idx = i * step
        slides_data.append({
            "title": slide_states[idx],
            "content": slide_states[idx + 1],
            "layout": slide_states[idx + 2],
            "insights": slide_states[idx + 3],
            "questions": slide_states[idx + 4],
            "feedback": slide_states[idx + 5],
            "preview": slide_states[idx + 6],
            "approved": slide_states[idx + 7]
        })
        
    session_payload = {
        "session_name": session_name,
        "slide_count": int(count),
        "outline": outline,
        "slides": slides_data
    }
    
    file_path = os.path.join(SESSIONS_DIR, f"{session_name.strip()}.json")
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(session_payload, f, indent=2)
        
    return f"💾 Session state saved safely to local storage: {file_path}"


def load_user_session(session_name):
    """Locates target local JSON state files and hydrates values back into active Gradio tab components."""
    file_path = os.path.join(SESSIONS_DIR, f"{session_name.strip()}.json")
    if not os.path.exists(file_path):
        return ["❌ Session target name not found!"] + [gr.update() for _ in range(2 + MAX_SLIDES * 9)]
        
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    count = data["slide_count"]
    updates = [
        f"📂 Successfully loaded session state profile: '{session_name}'",
        gr.update(value=count),
        gr.update(value=data["outline"]),
        gr.update(visible=True)
    ]
    
    for i in range(MAX_SLIDES):
        if i < count:
            s = data["slides"][i]
            updates.extend([
                gr.update(visible=True),  # 1. slide_tab visibility
                gr.update(value=s["title"]),
                gr.update(value=s["content"]),
                gr.update(value=s["layout"]),
                gr.update(value=s["insights"]),
                gr.update(value=s["questions"]),
                gr.update(value=s["feedback"]),
                gr.update(value=s["preview"]),
                gr.update(value=s["approved"])
            ])
        else:
            updates.extend([gr.update(visible=False)] + [gr.update() for _ in range(8)])
            
    return updates


def compile_presentation_file(session_name, *slide_states):
    """Processes all active locked/approved slides using python-pptx and builds an output file package."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    
    OUTPUT_DIR = "generated_presentations"
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    step = 8
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts
    
    approved_slides_count = 0
    
    for i in range(MAX_SLIDES):
        idx = i * step
        title_text = slide_states[idx]
        content_text = slide_states[idx + 1]
        layout_choice = slide_states[idx + 2]
        is_approved = slide_states[idx + 7]
        
        if is_approved:
            approved_slides_count += 1
            slide = prs.slides.add_slide(blank_slide_layout)
            
            dark_bg = RGBColor(0x1F, 0x29, 0x37)
            accent_blue = RGBColor(0x25, 0x63, 0xEB)
            text_color = RGBColor(0x37, 0x41, 0x51)
            
            if layout_choice == "Minimalist Title Centered":
                txBox = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9.333), Inches(2.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs
                p.text = title_text
                p.font.size = Pt(48)
                p.font.bold = True
                p.font.color.rgb = dark_bg
                if content_text:
                    p2 = tf.add_paragraph()
                    p2.text = content_text
                    p2.font.size = Pt(20)
                    p2.font.color.rgb = accent_blue
                    
            elif layout_choice == "Split Text & Image":
                txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1), Inches(6), Inches(5.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs
                p.text = title_text
                p.font.size = Pt(36)
                p.font.bold = True
                p2 = tf.add_paragraph()
                p2.text = content_text
                p2.font.size = Pt(18)
                p2.font.color.rgb = text_color
                shape = slide.shapes.add_shape(1, Inches(7.5), Inches(1.5), Inches(5), Inches(4.5))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
                shape.line.color.rgb = accent_blue
            else:
                txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11.333), Inches(5.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs
                p.text = title_text
                p.font.size = Pt(40)
                p.font.bold = True
                p2 = tf.add_paragraph()
                p2.text = f"[{layout_choice}]\n\n{content_text}"
                p2.font.size = Pt(20)
                
    if approved_slides_count == 0:
        return None, "⚠️ Compilation stalled: You must check 'Lock & Approve' on at least one slide tab."
        
    filename = f"{session_name.strip() if session_name.strip() else 'Presentation'}_compiled.pptx"
    output_path = os.path.join(OUTPUT_DIR, filename)
    prs.save(output_path)
    
    return output_path, f"🚀 **Compilation successful!** Packaged {approved_slides_count} slides."

def on_thumbnail_select(evt: gr.SelectData):
    return evt.value

def trigger_comfyui_generation(title, content, layout, feedback):
    yield gr.update(value="⏳ Rendering via ComfyUI Pipeline..."), gr.update(interactive=False)
    time.sleep(1.5)
    placeholder_image = "https://images.unsplash.com/photo-1618005182384-a83a8bd57fbe?w=500"
    yield gr.update(value=placeholder_image), gr.update(interactive=True, value="🎨 Re-Generate Assets")
