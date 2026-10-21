"""Gradio slide workspace helpers, persistence, and PowerPoint compilation."""

import json
import os
import time
from pathlib import Path

import gradio as gr
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

MAX_SLIDES = 25
SESSIONS_DIR = "presentation_sessions"
OUTPUT_DIR = "generated_presentations"
PPT_DIR = "generated_presentations/ppt2"

for directory in (SESSIONS_DIR, OUTPUT_DIR, PPT_DIR):
    os.makedirs(directory, exist_ok=True)

LAYOUT_TEMPLATES = {
    "Minimalist Title Centered": "https://images.unsplash.com/photo-1626785774573-4b799315345d?w=150&q=80",
    "Split Text & Image": "https://images.unsplash.com/photo-1618005182384-a83a8bd57fbe?w=150&q=80",
    "Three-Column Feature Grid": "https://images.unsplash.com/photo-1557683316-973673baf926?w=150&q=80",
    "Technical Diagram Grid": "https://images.unsplash.com/photo-1579546929518-9e396f3cc809?w=150&q=80",
    "Big Metric Callout": "https://images.unsplash.com/photo-1614850523459-c2f4c699c52e?w=150&q=80",
}
GALLERY_ITEMS = [(image, name) for name, image in LAYOUT_TEMPLATES.items()]


def _safe_slide_count(value):
    try:
        count = int(value)
    except (TypeError, ValueError):
        count = 1
    return max(1, min(MAX_SLIDES, count))


def _clean_outline_lines(outline_text):
    lines = []
    for line in (outline_text or "").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        # Remove optional numeric prefixes such as "1. " or "1) ".
        import re
        line = re.sub(r"^\s*\d+\s*[\.\)]\s*", "", line).strip()
        if line:
            lines.append(line)
    return lines


def process_outline_and_suggest_layouts(slide_count, outline_text):
    """Return workspace visibility plus 5 updates per slide slot."""
    count = _safe_slide_count(slide_count)
    titles = _clean_outline_lines(outline_text)
    updates = [gr.update(visible=True)]

    for i in range(MAX_SLIDES):
        if i < count:
            title = titles[i] if i < len(titles) else f"Slide {i + 1} Title"
            lower = title.lower()
            layout = "Split Text & Image"
            if "architecture" in lower or "technical" in lower or "workflow" in lower:
                layout = "Technical Diagram Grid"
            elif "introduction" in lower or "conclusion" in lower or "title" in lower:
                layout = "Minimalist Title Centered"

            updates.extend([
                gr.update(visible=True),        # accordion/group
                gr.update(value=title),          # title
                gr.update(value=""),             # content; filled by LLM when available
                gr.update(value=layout),         # layout
                gr.update(value=""),             # feedback
            ])
        else:
            updates.extend([
                gr.update(visible=False),
                gr.update(),
                gr.update(),
                gr.update(),
                gr.update(),
            ])
    return updates


def on_thumbnail_select(evt: gr.SelectData):
    """Return the selected gallery caption."""
    value = evt.value
    if isinstance(value, (tuple, list)) and len(value) > 1:
        return value[1]
    return str(value)


def trigger_comfyui_generation(title, content, layout, feedback):
    """Compatibility wrapper for the render button; real generation is in comfyui.py."""
    return gr.update(value=None), gr.update(interactive=True, value="🎨 Render slide")


def _session_path(session_name):
    name = Path((session_name or "").strip()).name
    if not name:
        return None
    if not name.endswith(".json"):
        name += ".json"
    return os.path.join(SESSIONS_DIR, name)


def save_user_session(session_name, count, outline, *slide_states):
    """Save title, content, layout and feedback for each of 25 slides."""
    path = _session_path(session_name)
    if path is None:
        return "❌ Please enter a session name."

    slides = []
    step = 4
    for i in range(MAX_SLIDES):
        idx = i * step
        values = list(slide_states[idx:idx + step])
        values += [""] * (step - len(values))
        title, content, layout, feedback = values
        slides.append({
            "title": title or "",
            "content": content or "",
            "layout": layout or "",
            "feedback": feedback or "",
        })

    payload = {
        "session_name": Path(path).stem,
        "slide_count": _safe_slide_count(count),
        "outline": outline or "",
        "slides": slides,
    }
    with open(path, "w", encoding="utf-8") as file:
        json.dump(payload, file, indent=2, ensure_ascii=False)
    return f"💾 Session saved: `{path}`"


def load_user_session(session_name):
    """Return status, count, outline, workspace visibility and 5 updates per slot."""
    path = _session_path(session_name)
    if path is None or not os.path.isfile(path):
        return ["❌ Session not found."] + [gr.update() for _ in range(3 + MAX_SLIDES * 5)]

    try:
        with open(path, "r", encoding="utf-8") as file:
            data = json.load(file)
    except (OSError, json.JSONDecodeError) as exc:
        return [f"❌ Could not load session: {exc}"] + [gr.update() for _ in range(3 + MAX_SLIDES * 5)]

    count = _safe_slide_count(data.get("slide_count", 1))
    updates = [
        f"📂 Loaded session: **{data.get('session_name', Path(path).stem)}**",
        gr.update(value=count),
        gr.update(value=data.get("outline", "")),
        gr.update(visible=True),
    ]
    slides = data.get("slides", [])
    for i in range(MAX_SLIDES):
        slide = slides[i] if i < len(slides) else {}
        visible = i < count
        updates.extend([
            gr.update(visible=visible),
            gr.update(value=slide.get("title", "")),
            gr.update(value=slide.get("content", "")),
            gr.update(value=slide.get("layout", "")),
            gr.update(value=slide.get("feedback", "")),
        ])
    return updates


def compile_presentation_file(session_name, *slide_states):
    """Compile slides with non-empty titles into a widescreen PowerPoint."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    step = 4
    approved = 0

    for i in range(MAX_SLIDES):
        idx = i * step
        values = list(slide_states[idx:idx + step])
        values += [""] * (step - len(values))
        title, content, layout_choice, _feedback = values
        title = str(title or "").strip()
        content = str(content or "").strip()
        layout_choice = str(layout_choice or "Split Text & Image")
        if not title:
            continue

        approved += 1
        slide = prs.slides.add_slide(blank_layout)
        navy = RGBColor(31, 41, 55)
        blue = RGBColor(37, 99, 235)
        dark_text = RGBColor(55, 65, 81)

        if layout_choice == "Minimalist Title Centered":
            box = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.2))
            box.text_frame.word_wrap = True
            p = box.text_frame.paragraphs[0]
            p.text = title
            p.alignment = 1
            p.font.size = Pt(42)
            p.font.bold = True
            p.font.color.rgb = navy
            if content:
                p2 = box.text_frame.add_paragraph()
                p2.text = content
                p2.alignment = 1
                p2.font.size = Pt(18)
                p2.font.color.rgb = blue
        elif layout_choice == "Split Text & Image":
            box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(6.0), Inches(5.8))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(34)
            p.font.bold = True
            p.font.color.rgb = navy
            if content:
                p2 = tf.add_paragraph()
                p2.text = content
                p2.font.size = Pt(17)
                p2.font.color.rgb = dark_text
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(7.3), Inches(1.3), Inches(5.2), Inches(4.7)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(229, 231, 235)
            shape.line.color.rgb = blue
        elif layout_choice == "Big Metric Callout":
            box = slide.shapes.add_textbox(Inches(1), Inches(1.4), Inches(11.3), Inches(4.8))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(38)
            p.font.bold = True
            p.font.color.rgb = blue
            if content:
                p2 = tf.add_paragraph()
                p2.text = content
                p2.font.size = Pt(20)
                p2.font.color.rgb = dark_text
        else:
            box = slide.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(11.5), Inches(5.8))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(34)
            p.font.bold = True
            p.font.color.rgb = navy
            if content:
                p2 = tf.add_paragraph()
                p2.text = content
                p2.font.size = Pt(18)
                p2.font.color.rgb = dark_text

    if approved == 0:
        return None, "⚠️ Add a title to at least one slide before compiling."

    safe_name = Path((session_name or "Presentation").strip() or "Presentation").stem
    output_path = os.path.join(OUTPUT_DIR, f"{safe_name}_compiled.pptx")
    prs.save(output_path)
    return output_path, f"🚀 Compilation successful: **{approved} slides** written to the PowerPoint file."
