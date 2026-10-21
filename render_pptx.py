"""Render a Deck to an editable .pptx (real placeholders, speaker notes, 16:9)
and optionally export it to PDF through LibreOffice."""
import shutil
import subprocess
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from schema import Deck, Slide

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
INK = RGBColor(0x22, 0x28, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xC9, 0xD3, 0xE6)
FONT = "Calibri"

# Default template layout indexes
L_TITLE, L_CONTENT, L_SECTION, L_TWO = 0, 1, 2, 3


def _widen_template(prs: Presentation) -> None:
    """The default template is 4:3. Switch to 16:9 and stretch master/layout
    placeholders so 'Reset layout' in PowerPoint stays sensible."""
    old_w = prs.slide_width
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    factor = prs.slide_width / old_w
    holders = list(prs.slide_master.shapes)
    for layout in prs.slide_layouts:
        holders += list(layout.shapes)
    for shp in holders:
        if shp.left is not None and shp.width is not None:
            shp.left = int(shp.left * factor)
            shp.width = int(shp.width * factor)


def _place(shape, left, top, width, height) -> None:
    shape.left, shape.top, shape.width, shape.height = (
        Inches(left), Inches(top), Inches(width), Inches(height),
    )


def _style_runs(paragraph, size: int, color: RGBColor, bold: bool = False) -> None:
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def _set_text(shape, text: str, size: int, color: RGBColor, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    _style_runs(p, size, color, bold)


def _set_bullets(shape, bullets: list[str], size: int) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(10)
        _style_runs(p, size, INK)


def _body_size(texts: list[str], columns: int = 1) -> int:
    """Cheap overflow guard: shrink the font as content grows."""
    chars = sum(len(t) for t in texts) + 12 * len(texts)
    score = chars if columns == 1 else chars / 0.55
    if score < 260:
        return 28
    if score < 380:
        return 24
    if score < 520:
        return 20
    return 17


def _accent_bar(slide) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.4), Inches(0.25), Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    bar.name = "Accent bar"


def _dark_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def _notes(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _title_slide(prs, deck: Deck) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    _dark_background(slide)
    title, subtitle = slide.placeholders[0], slide.placeholders[1]
    _place(title, 0.9, 2.3, 11.5, 1.8)
    _place(subtitle, 0.9, 4.3, 11.5, 1.0)
    _set_text(title, deck.title, 40 if len(deck.title) > 50 else 48, WHITE, bold=True,
              anchor=MSO_ANCHOR.BOTTOM)
    _set_text(subtitle, deck.subtitle, 22, MUTED, anchor=MSO_ANCHOR.TOP)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.15), Inches(1.6), Inches(0.06))
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()


def _section_slide(prs, s: Slide) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    _dark_background(slide)
    title, body = slide.placeholders[0], slide.placeholders[1]
    _place(title, 0.9, 2.5, 11.5, 1.6)
    _place(body, 0.9, 4.2, 11.5, 1.2)
    _set_text(title, s.title, 40, WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _set_text(body, s.bullets[0] if s.bullets else "", 22, MUTED, anchor=MSO_ANCHOR.TOP)
    _notes(slide, s.notes)


def _content_slide(prs, s: Slide) -> None:
    two = s.layout == "two_column" and s.right_bullets
    slide = prs.slides.add_slide(prs.slide_layouts[L_TWO if two else L_CONTENT])
    title = slide.placeholders[0]
    _place(title, 0.8, 0.4, 11.7, 1.1)
    _set_text(title, s.title, 30 if len(s.title) > 45 else 34, NAVY, bold=True)
    _accent_bar(slide)
    if two:
        left, right = slide.placeholders[1], slide.placeholders[2]
        size = _body_size(s.bullets + s.right_bullets, columns=2)
        _place(left, 0.8, 1.8, 5.7, 5.0)
        _place(right, 6.8, 1.8, 5.7, 5.0)
        _set_bullets(left, s.bullets, size)
        _set_bullets(right, s.right_bullets, size)
    else:
        body = slide.placeholders[1]
        _place(body, 0.8, 1.8, 11.7, 5.0)
        _set_bullets(body, s.bullets, _body_size(s.bullets))
    _notes(slide, s.notes)


def render_pptx(deck: Deck, path: str | Path) -> Path:
    prs = Presentation()
    _widen_template(prs)
    _title_slide(prs, deck)
    for s in deck.slides:
        if s.layout == "section":
            _section_slide(prs, s)
        else:
            _content_slide(prs, s)
    path = Path(path)
    prs.save(path)
    return path


def export_pdf(pptx_path: str | Path) -> Optional[Path]:
    """Convert with LibreOffice if installed. Returns the PDF path or None."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    pptx_path = Path(pptx_path)
    subprocess.run(
        [exe, "--headless", "--convert-to", "pdf", "--outdir", str(pptx_path.parent), str(pptx_path)],
        check=True, capture_output=True, timeout=180,
    )
    pdf = pptx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None
