"""Verifies that the batch flow keeps the LLM-structured content through the compile step."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import user_context as uc
import gemini_code_fixed as g
from layout_catalogue import get_template
import test_layout_pipeline as T

fake = T.FakeOllama()
# the batch handler uses the module level client - swap in the mock for a fast,
# deterministic run
uc.ollama = fake
g.ollama = fake

# a deck with three slides, each with a different template
plan = [("L01", "Opening the review"), ("L31", "The numbers"), ("L08", "Two options")]
states = []
for i, (lid, title) in enumerate(plan):
    states.extend([title, "raw content for the slide", get_template(lid).label, ""])
states.extend([""] * ((25 - len(plan)) * 4))

# the batch handler is what the 'Generate all slides' button calls
out = uc._generate_all_slides("BatchTest", "gpt-oss", *states)
path, status = out[0], out[1]
print("deck      :", path, os.path.getsize(path), "bytes")
print("status    :", status.encode("ascii", "replace").decode().replace("\n", " ")[:100])
print("structured:", {i: (g.SLIDE_TEMPLATES.get(i), sorted(g.SLIDE_STRUCTURED[i].keys()))
                      for i in range(3)})

# the compiled deck must carry the LLM content, not the heuristic fallback
# (the slide title stays the user's own title by design - the model's own wording is
#  used for every other field)
from pptx import Presentation
prs = Presentation(path)
for i, slide in enumerate(prs.slides):
    texts = [sh.text_frame.text for sh in slide.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    print(f"slide {i+1}: {len(slide.shapes)} shapes | " + " || ".join(t.replace(chr(10), " / ")
                                                                    for t in texts[:3]))

deck_text = " ".join(sh.text_frame.text for slide in prs.slides for sh in slide.shapes
                     if sh.has_text_frame)
assert "OVERVIEW" in deck_text, "the LLM structured content did not reach the compiled deck"
assert "A restructured subtitle" in deck_text, "the LLM subtitle did not reach the compiled deck"
assert "Opening the review" in deck_text, "the user title was not preserved"
print("\nBATCH FLOW VERIFIED")
